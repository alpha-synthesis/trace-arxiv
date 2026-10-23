#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Human audit of the R2 news contexts: sampling and workbook.

Draws 100 of the 4,340 contexts (cache/contexts_v2.1) with a fixed seed. In
each sampled context it lists every item that mentions a market, index,
price or percentage term (where an outcome report would appear) plus three
further items drawn at random from the rest, which measure what the filter
misses. The auditor marks each item blind to the window's return.

Output: audit/human_context_audit.xlsx (sheets Instructions, Items, Summary)
and audit/human_context_audit_sample.json (the sample, for scoring).

Usage: python build_human_audit.py
"""
from __future__ import annotations

import glob
import json
import re
import sys
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import ORIG  # noqa: E402

SEED = 20261008
N_CONTEXTS = 100
N_RANDOM_ITEMS = 3
CTX = ORIG / "cache" / "contexts_v2.1"
DST = Path(__file__).resolve().parents[1] / "audit"

FLAG = re.compile(
    r"\b(stocks?|shares?|share price|index|indices|indexes|Dow|S&P|Nasdaq|FTSE|Nikkei|DAX|Hang Seng|"
    r"markets?|Wall Street|equit\w+|investors?|trading|traders?|sell-?off|rall\w+|plung\w+|surg\w+|"
    r"tumbl\w+|soar\w+|slump\w+|crash\w*|price[sd]?|percent|points?|valuation|bear|bull\w*)\b|%",
    re.I)
ITEM = re.compile(r"^\[(\d+)\] \((\d{4}-\d{2}-\d{2})\) ", re.M)


def _items(context):
    starts = [(m.start(), m.end(), int(m.group(1)), m.group(2)) for m in ITEM.finditer(context)]
    out = []
    for i, (s, e, n, d) in enumerate(starts):
        end = starts[i + 1][0] if i + 1 < len(starts) else len(context)
        text = "\n".join(l for l in context[e:end].strip().splitlines() if not l.startswith("──")).strip()
        out.append({"item": n, "date": d, "text": text})
    return out


def sample():
    ctx = []
    for f in sorted(glob.glob(str(CTX / "[A-Z]*.json"))):
        j = json.load(open(f))
        if "windows" in j:
            ctx += [(j["ticker"], w) for w in j["windows"]]
    rng = np.random.RandomState(SEED)
    picked = sorted(rng.choice(len(ctx), N_CONTEXTS, replace=False))
    rows = []
    for cid, k in enumerate(picked, 1):
        t, w = ctx[k]
        items = [x for x in _items(w["context"]) if x["text"]]
        flagged = [x for x in items if FLAG.search(x["text"])]
        rest = [x for x in items if not FLAG.search(x["text"])]
        extra = [rest[i] for i in sorted(rng.choice(len(rest), min(N_RANDOM_ITEMS, len(rest)), replace=False))] if rest else []
        for sel, xs in (("flagged", flagged), ("random", extra)):
            rows += [{"context": cid, "stock": t, "window_start": w["window_start"], "window_end": w["window_end"],
                      "context_sha1": w["context_sha1"], "selection": sel, **x} for x in xs]
    return len(ctx), rows


INSTRUCTIONS = [
    ("Purpose", "Check whether the news contexts given to the LLM (price-blind arm R2) reveal the outcome the labels "
                "are meant to explain. 100 of the 4,340 contexts were drawn at random (seed 20261008). Each context "
                "covers one stock and one two-week window."),
    ("What to read", "Sheet Items lists, for each sampled context, every item mentioning a market, index, price or "
                     "percentage term (selection = flagged) and three further items drawn at random from the rest "
                     "(selection = random). Read each item; you do not need the window's return, and should not look it up."),
    ("Q1 equity outcome", "Y if the item reports how stock prices moved during the window: the stock's own price, a stock "
                          "index (Dow, S&P 500, Nasdaq, FTSE, Nikkei...), 'stocks', 'Wall Street' or 'the market' rising "
                          "or falling, a sell-off, a crash, a rally, trading halts or circuit breakers. Otherwise N."),
    ("Q2 other price move", "Y if the item reports a non-equity price move (oil, gold, currencies, crypto, bond yields). "
                            "These are kept by design and are reported separately. Otherwise N."),
    ("Q3 after the window", "Y if the item describes an event dated after the window end (look-ahead). Otherwise N."),
    ("Not a leak", "Company or economic news without a price move (earnings released, layoffs, a rate decision, "
                   "unemployment data) is information the labels may use: mark N on Q1 to Q3."),
    ("Where to write", "Only the yellow columns Q1, Q2, Q3 (Y or N, drop-down) and Notes. Leave the other columns as they are. "
                       "Sheet Summary updates by itself."),
    ("Time", "About 2 to 3 hours. Save the file under the same name when done."),
]
EXAMPLE = ("Example (not part of the sample)", "AAPL", "2020-03-01", "2020-03-15", "flagged", 12, "2020-03-09",
           "The Dow Jones Industrial Average falls 2,014 points, its largest point drop in history.", "Y", "N", "N",
           "Index fall within the window")


def workbook(n_total, rows, path):
    wb = Workbook()
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    fill = PatternFill("solid", fgColor="FFFF00")

    ws = wb.active
    ws.title = "Instructions"
    ws["A1"], ws["A1"].font = "Human audit of the R2 news contexts", Font(name="Arial", size=12, bold=True)
    for i, (k, v) in enumerate(INSTRUCTIONS, 3):
        ws.cell(i, 1, k).font = bold
        c = ws.cell(i, 2, v)
        c.font, c.alignment = font, wrap
    r0 = len(INSTRUCTIONS) + 4
    hdr = ["", "Stock", "Window start", "Window end", "Selection", "Item", "Item date", "Text", "Q1", "Q2", "Q3", "Notes"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r0, j, h).font = bold
    for j, v in enumerate(EXAMPLE, 1):
        c = ws.cell(r0 + 1, j, v)
        c.font, c.alignment = font, wrap
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 100
    ws.column_dimensions["H"].width = 60

    it = wb.create_sheet("Items")
    cols = ["Context", "Stock", "Window start", "Window end", "Selection", "Item", "Item date", "Text",
            "Q1 equity outcome", "Q2 other price move", "Q3 after window", "Notes"]
    for j, h in enumerate(cols, 1):
        it.cell(1, j, h).font = bold
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    it.add_data_validation(dv)
    for i, r in enumerate(rows, 2):
        vals = [r["context"], r["stock"], r["window_start"], r["window_end"], r["selection"], r["item"], r["date"], r["text"]]
        for j, v in enumerate(vals, 1):
            c = it.cell(i, j, v)
            c.font, c.alignment = font, wrap
        for j in (9, 10, 11, 12):
            c = it.cell(i, j)
            c.fill, c.font, c.alignment = fill, font, wrap
        dv.add(f"I{i}:K{i}")
    for col, wdt in zip("ABCDEFGHIJKL", (8, 7, 11, 11, 9, 5, 11, 80, 10, 10, 10, 30)):
        it.column_dimensions[col].width = wdt
    it.freeze_panes = "A2"
    last = len(rows) + 1

    sm = wb.create_sheet("Summary")
    sm["A1"], sm["A1"].font = "Summary (updates as Items is filled)", Font(name="Arial", size=12, bold=True)
    lines = [
        ("Contexts in the population", n_total, "Count of R2 contexts (28 stocks x 155 windows)"),
        ("Contexts sampled", N_CONTEXTS, "Random draw, seed 20261008"),
        ("Items listed", f"=COUNTA(Items!H2:H{last})", ""),
        ("Items flagged by the market/price filter", f'=COUNTIF(Items!E2:E{last},"flagged")', ""),
        ("Random unflagged items", f'=COUNTIF(Items!E2:E{last},"random")', ""),
        ("Items marked so far (Q1 filled)", f"=COUNTA(Items!I2:I{last})", "Equal to items listed when done"),
        ("Q1 = Y (equity outcome)", f'=COUNTIF(Items!I2:I{last},"Y")', ""),
        ("Q1 = Y among random unflagged items", f'=COUNTIFS(Items!I2:I{last},"Y",Items!E2:E{last},"random")', "Measures what the filter misses"),
        ("Q2 = Y (other price move, kept by design)", f'=COUNTIF(Items!J2:J{last},"Y")', ""),
        ("Q3 = Y (after the window)", f'=COUNTIF(Items!K2:K{last},"Y")', ""),
        ("Contexts with any Q1 or Q3 = Y", f'=SUMPRODUCT((MATCH(Items!A2:A{last},Items!A2:A{last},0)=ROW(Items!A2:A{last})-1)*'
                                            f'(COUNTIFS(Items!A2:A{last},Items!A2:A{last},Items!I2:I{last},"Y")+'
                                            f'COUNTIFS(Items!A2:A{last},Items!A2:A{last},Items!K2:K{last},"Y")>0))', ""),
        ("95% upper bound, share of contexts with a leak", "=BETAINV(0.95,B13+1,B4-B13)",
         "One-sided Clopper-Pearson; about 3/n when none is found"),
    ]
    for i, (k, v, note) in enumerate(lines, 3):
        sm.cell(i, 1, k).font = bold
        c = sm.cell(i, 2, v)
        c.font = font
        sm.cell(i, 3, note).font = font
    sm["B14"].number_format = "0.0%"
    sm.column_dimensions["A"].width, sm.column_dimensions["B"].width, sm.column_dimensions["C"].width = 48, 12, 50
    wb.save(path)


def main():
    DST.mkdir(parents=True, exist_ok=True)
    n_total, rows = sample()
    json.dump({"seed": SEED, "n_contexts": N_CONTEXTS, "population": n_total, "rows": rows},
              open(DST / "human_context_audit_sample.json", "w"), indent=1)
    workbook(n_total, rows, DST / "human_context_audit.xlsx")
    flagged = sum(r["selection"] == "flagged" for r in rows)
    words = sum(len(r["text"].split()) for r in rows)
    print(f"{len(rows)} items ({flagged} flagged, {len(rows) - flagged} random), {words} words")


if __name__ == "__main__":
    main()
